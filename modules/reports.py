from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from docx import Document
from docx.shared import Inches
import io, os, time

REPORT_DIR = Path(__file__).resolve().parent.parent / "reports"
REPORT_DIR.mkdir(exist_ok=True)

def make_pdf(product, category, stats, rows, image=None):
    path = REPORT_DIR / f"LabelLens_{int(time.time())}.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=35,leftMargin=35,topMargin=35,bottomMargin=35)
    styles = getSampleStyleSheet()
    story = [Paragraph("LabelLens — Compliance Inspection Report", styles["Title"]),
             Paragraph(f"Product: {product or 'Unknown'} | Category: {category.title()}", styles["Normal"]),
             Spacer(1, 12),
             Paragraph(f"Compliance Score: {stats['score']}/100", styles["Heading2"])]
    if image:
        bio = io.BytesIO()
        image.save(bio, format="JPEG")
        bio.seek(0)
        story += [RLImage(bio, width=4.8*inch, height=3.2*inch), Spacer(1,8)]
    data = [["Rule","Field","Status","Evidence"]]
    for r in rows:
        data.append([r["Rule"], r["Field"], r["Status"], r["Evidence"][:80]])
    table = Table(data, colWidths=[45,110,105,245])
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#102A43")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("GRID",(0,0),(-1,-1),0.4,colors.grey),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("FONTSIZE",(0,0),(-1,-1),7),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white, colors.HexColor("#F4F7FA")])
    ]))
    story += [table, Spacer(1,10), Paragraph("Prototype note: results are decision-support output and should be verified by an authorized reviewer for enforcement use.", styles["Italic"])]
    doc.build(story)
    return path

def make_docx(product, category, stats, rows, image=None):
    path = REPORT_DIR / f"LabelLens_{int(time.time())}.docx"
    d = Document()
    d.add_heading("LabelLens — Compliance Inspection Report", 0)
    d.add_paragraph(f"Product: {product or 'Unknown'} | Category: {category.title()}")
    d.add_heading(f"Compliance Score: {stats['score']}/100", 1)
    if image:
        tmp = REPORT_DIR / f"_tmp_{int(time.time())}.jpg"
        image.save(tmp)
        d.add_picture(str(tmp), width=Inches(5.5))
        tmp.unlink(missing_ok=True)
    table = d.add_table(rows=1, cols=4)
    for i,h in enumerate(["Rule","Field","Status","Evidence"]): table.rows[0].cells[i].text=h
    for r in rows:
        cells=table.add_row().cells
        for i,k in enumerate(["Rule","Field","Status","Evidence"]): cells[i].text=str(r[k])
    d.add_paragraph("Prototype note: results should be verified by an authorized reviewer for enforcement use.")
    d.save(path)
    return path
